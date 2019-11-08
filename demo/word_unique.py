import xlwt
import win32com.client  # 系统客户端包
import docx

speaker = win32com.client.Dispatch("SAPI.SPVOICE")  # 系统接口

doc_file_path = r'C:\Users\Administrator\Desktop\词汇复习手册.docx'
excel_path = r'C:\Users\Administrator\Desktop\word_list.xls'

doc_file = docx.Document(doc_file_path)

# read file and save as dictory
dict = {}
characteristc = set()
for para in doc_file.paragraphs:
    comment = para.text.replace(para.text.split('  ')[0], '').strip()
    word = para.text.split('  ')[0]
    if word in dict.keys():
        old_comment = dict[word]
        if len(old_comment) < len(comment):
            dict[word] = comment
            print(word,'old_comment',old_comment,'comment',comment)
        else:
            continue
    else:
        dict[word] = comment
    if '  ' in word or '  ' in comment:
        print(word, comment)
    for i in range(0,comment.count('.'),2):
        characteristc.add(comment.split('.')[i].split(' ')[-1])
    if '.' not  in comment:
        print(word,dict[word])
    # speaker.Speak(word)
    # speaker.Speak(comment.replace('adv','ad').replace('adj','ad').replace('pron','代词').replace('conj','连词').replace('v','动词'))

print(characteristc)

# write as excel
workbook = xlwt.Workbook(excel_path)
worksheet = workbook.add_sheet('Word_List', cell_overwrite_ok='True')

i = 1
worksheet.write(0, 0, '单词')
worksheet.write(0, 1, '注释')

for word_comment in dict:

    worksheet.write(i, 0, word_comment)
    worksheet.write(i, 1, dict[word_comment])

    i += 1
workbook.save(excel_path)
